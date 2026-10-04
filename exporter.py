from datetime import datetime

from loguru import logger
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from scraper import Product

PRICE_COL = 3
RATING_COL = 4
URL_COL = 7


def build_output_path(output_dir):
    output_dir.mkdir(parents=True, exist_ok=True)
    return output_dir / datetime.now().strftime(config.OUTPUT_FILENAME_PATTERN)


def export_products_to_excel(products, output_dir=config.OUTPUT_DIR):
    wb = Workbook()
    ws = wb.active
    ws.title = config.EXCEL_SHEET_NAME

    ws.append(config.EXCEL_COLUMNS)
    header_fill = PatternFill("solid", start_color=config.HEADER_FILL_COLOR)
    header_font = Font(bold=True, color=config.HEADER_FONT_COLOR)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    rows = [p.to_row() for p in products]
    for row in rows:
        ws.append(row)

    for row in ws.iter_rows(min_row=2):
        row[PRICE_COL - 1].number_format = config.PRICE_NUMBER_FORMAT
        row[RATING_COL - 1].number_format = "0"
        url_cell = row[URL_COL - 1]
        if url_cell.value:
            url_cell.hyperlink = url_cell.value
            url_cell.style = "Hyperlink"

    for col_idx in range(1, len(config.EXCEL_COLUMNS) + 1):
        values = [len(str(row[col_idx - 1])) for row in rows]
        width = max(values + [len(config.EXCEL_COLUMNS[col_idx - 1])]) + 2
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width, 60)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    path = build_output_path(output_dir)
    wb.save(path)
    logger.success("saved: {}", path)
    return path
